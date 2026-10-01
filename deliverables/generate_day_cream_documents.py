"""Generate the data-locked PDF and PowerPoint R&D pack for the Daylight Day Cream.
Both documents are built from the single FORMULA list below so ingredient, % and gram values cannot drift.
Run: python3 deliverables/generate_day_cream_documents.py
"""
from __future__ import annotations

from pathlib import Path
from math import pi
from xml.etree import ElementTree as ET

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor, Color
from reportlab.pdfbase.pdfmetrics import stringWidth

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

OUT = Path(__file__).parent
PDF_PATH = OUT / "Daylight_Day_Cream_RnD_Formula_Tutorial.pdf"
PPTX_PATH = OUT / "Daylight_Day_Cream_RnD_Presentation.pptx"

# ---------------------------------------------------------------------------
# Canonical formula data. Every appearance of the master formula is generated
# from this list. This is intentionally fragrance-free.
# ---------------------------------------------------------------------------
FORMULA = [
    ("A", "Purified / deionized water", "Main carrier for a clean, light emulsion.", 43.69, "A water"),
    ("A", "Glycerin", "Immediate humectant; pulls water into the outer skin.", 4.00, "A water"),
    ("A", "1,3-Propanediol", "Humectant and slip aid; helps a less tacky feel.", 3.00, "A water"),
    ("A", "Betaine (trimethylglycine)", "Osmolyte humectant for comfortable hydration.", 1.00, "A water"),
    ("A", "Allantoin", "Skin-conditioning powder; supports a smoother feel.", 0.20, "A water"),
    ("A", "Disodium EDTA", "Chelator; supports preservative and formula stability.", 0.10, "A water"),
    ("A", "Xanthan gum (transparent grade)", "Low-dose thickener; helps keep the emulsion stable.", 0.15, "A water"),
    ("B", "Caprylic/Capric Triglyceride", "Light emollient for softness without a waxy film.", 4.00, "B oil"),
    ("B", "Dicaprylyl Carbonate", "Fast-spreading dry emollient for quick absorption.", 3.00, "B oil"),
    ("B", "Squalane", "Stable, skin-compatible emollient; supports softness and barrier feel.", 2.00, "B oil"),
    ("B", "Olivem 1000 (Cetearyl Olivate and Sorbitan Olivate)", "Primary emulsifier; binds water and oils into cream.", 4.50, "B oil"),
    ("B", "Cetearyl alcohol", "Co-emulsifier; gives cushion, body and stability.", 0.70, "B oil"),
    ("B", "Mixed tocopherol (min. 70% active)", "Oil-phase antioxidant; protects the emollient phase.", 0.20, "B oil"),
    ("C", "Purified / deionized water", "Cool-down carrier for dissolved actives.", 25.00, "C water"),
    ("C", "Niacinamide", "Brightening/barrier-support active; selected at a studied 4% level.", 4.00, "C cool-down"),
    ("C", "N-Acetyl Glucosamine (NAG)", "Tone-evenness support; pairs well with niacinamide.", 2.00, "C cool-down"),
    ("C", "D-Panthenol (min. 98% powder)", "Humectant/barrier-support active for skin comfort.", 1.00, "C cool-down"),
    ("C", "Zinc PCA", "Blemish-prone/oil-balance support; not an acne medicine.", 0.30, "C cool-down"),
    ("C", "Trisodium citrate dihydrate", "Part of a gentle citrate buffer for the target pH.", 0.11, "C cool-down"),
    ("C", "Citric acid, anhydrous", "Part of the citrate buffer; supports pH control.", 0.05, "C cool-down"),
    ("C", "Phenoxyethanol (and) Ethylhexylglycerin", "Broad preservative blend (e.g., Euxyl PE 9010).", 1.00, "C cool-down"),
]

assert abs(sum(x[3] for x in FORMULA) - 100.0) < 0.00001, sum(x[3] for x in FORMULA)
PH_TARGET = "5.2-5.6 (aim 5.4 at 25 C)"
HEAT_RANGE = "70-75 C"
COOL_ADD = "below 40 C"
PRESERVATIVE_LIMIT = "1.00% blend; use a blend with phenoxyethanol within its supplier/regulatory limit"

# Normalized visual aliases, used only for display.
PHASES = {
    "A": "Water phase",
    "B": "Oil phase",
    "C": "Cool-down actives + preservation",
}

# Palette (premium spa / cosmetic lab).
INK = HexColor("#172A24")
FOREST = HexColor("#246653")
SAGE = HexColor("#DCE9DE")
MINT = HexColor("#EEF5EF")
CREAM = HexColor("#FBF8F0")
GOLD = HexColor("#C49751")
PEACH = HexColor("#F3E2D1")
TEAL = HexColor("#6CB3A2")
MUTED = HexColor("#60706A")
WHITE = HexColor("#FFFFFF")
LINE = HexColor("#D6DED7")
RED = HexColor("#B25048")

# PPT colors as RGB tuples.
C = {
    "ink": (23, 42, 36), "forest": (36, 102, 83), "sage": (220, 233, 222),
    "mint": (238, 245, 239), "cream": (251, 248, 240), "gold": (196, 151, 81),
    "peach": (243, 226, 209), "teal": (108, 179, 162), "muted": (96, 112, 106),
    "white": (255, 255, 255), "line": (214, 222, 215), "red": (178, 80, 72),
}

# ---------------------------------------------------------------------------
# Shared text
# ---------------------------------------------------------------------------
TIMELINE = [
    ("DAY 1-4", "More hydrated, smoother and fresher-looking; light reflects more evenly from a moisturised surface. Not pigment removal or acne cure."),
    ("1 WEEK", "Less tightness/flaking and more comfortable skin are realistic if the formula suits the skin. Breakouts may be unchanged."),
    ("2-4 WEEKS", "A gradual improvement in overall tone appearance and post-blemish look may begin. Results differ by skin, routine and sun exposure."),
    ("6-12 WEEKS", "Meaningful change in uneven pigmentation is a longer-term target. Daily broad-spectrum sunscreen is essential; melasma or persistent acne needs a dermatologist."),
]

SOURCE_LINES = [
    "[1] Bissett DL et al. J Cosmet Dermatol. 2007;6:20-26. 2% NAG, with greater effect in a 2% NAG + 4% niacinamide study after 8 weeks. PMID 17348991. https://pubmed.ncbi.nlm.nih.gov/17348991/",
    "[2] Hakozaki T et al. Br J Dermatol. 2002;147:20-31. Topical niacinamide and melanosome transfer / pigmentation appearance. DOI 10.1046/j.1365-2133.2002.04834.x",
    "[3] American Academy of Dermatology. Dark spots: tinted iron-oxide sunscreen, broad spectrum, water resistant, SPF 30+; reapply every two hours. https://www.aad.org/public/everyday-care/skin-care-secrets/routine/fade-dark-spots",
    "[4] SCCS/1575/16 (2016). Phenoxyethanol considered safe as a preservative in cosmetics at a maximum 1.0%. https://ec.europa.eu/health/scientific_committees/consumer_safety/docs/sccs_o_195.pdf",
    "[5] FDA Cosmetic Microbiological Safety (2012). A preservative efficacy / challenge test is used to confirm preservation for the actual formula and pack. https://downloads.regulations.gov/FDA-2011-N-0770-0009/attachment_1.pdf",
]

SOURCING = [
    ("Niacinamide", "cosmetic/Ph.Eur grade, >=99%", "100 g", "ASES Chemical Works lists 99.5% / 100 g at about Rs 240; confirm COA and stock."),
    ("Zinc PCA", "cosmetic grade, water soluble", "30 g", "Cosmesi Global lists cosmetic Zinc PCA; 100 g quoted around Rs 750; confirm COA and stock."),
    ("NAG + D-Panthenol", "NAG >=98%; D-Panthenol >=98%", "25 g each", "Buy only with batch COA/SDS from a cosmetic raw-material supplier."),
    ("Olivem 1000 + Cetearyl alcohol", "correct INCI / cosmetic grade", "100 g each", "Euroasia Cosmetics and other Indian cosmetic raw suppliers list cream emulsifiers; verify INCI."),
    ("CCT, Dicaprylyl Carbonate, Squalane", "cosmetic grade", "100 mL each", "Ask for oxidation/identity documents and date of manufacture."),
    ("Euxyl PE 9010 or verified equivalent", "Phenoxyethanol (and) Ethylhexylglycerin", "30 g", "Use the supplier-prescribed concentration and check the exact INCI; do not substitute a medicine."),
    ("Water phase / support materials", "DI water, glycerin USP, 1,3-propanediol, betaine, allantoin, EDTA, xanthan, citrate salts", "25-100 g", "Cosmetic/lab grade, with COA. Do not use unknown food or pharmacy powders as substitutes."),
]

COST_ROWS = [
    ("100 g", "Rs 65-110", "Rs 40-90", "Rs 105-200"),
    ("500 g", "Rs 280-450", "Rs 200-450", "Rs 480-900"),
    ("1 kg", "Rs 520-800", "Rs 400-900", "Rs 920-1,700"),
]

# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------
W, H = A4
M = 38

def pdf_wrap(text: str, font: str, size: float, maxw: float):
    words = text.split()
    lines, line = [], ""
    for word in words:
        trial = word if not line else line + " " + word
        if stringWidth(trial, font, size) <= maxw:
            line = trial
        else:
            if line:
                lines.append(line)
            # Avoid overflow on a URL / long word by keeping it intact rather than clipping
            line = word
    if line:
        lines.append(line)
    return lines


def pdf_text(c, x, y, text, size=10, color=INK, font="Helvetica", maxw=None, leading=None, align="left"):
    c.setFont(font, size)
    c.setFillColor(color)
    leading = leading or size * 1.28
    lines = pdf_wrap(text, font, size, maxw) if maxw else str(text).split("\n")
    for line in lines:
        xx = x
        if align == "center":
            xx = x - stringWidth(line, font, size) / 2
        elif align == "right":
            xx = x - stringWidth(line, font, size)
        c.drawString(xx, y, line)
        y -= leading
    return y


def pdf_card(c, x, y_top, width, height, fill=WHITE, stroke=LINE, radius=12):
    c.setFillColor(fill)
    c.setStrokeColor(stroke)
    c.setLineWidth(0.7)
    c.roundRect(x, y_top-height, width, height, radius, stroke=1, fill=1)


def pdf_rule(c, x1, y, x2, color=LINE, width=0.7):
    c.setStrokeColor(color)
    c.setLineWidth(width)
    c.line(x1, y, x2, y)


def pdf_badge(c, x, y_top, label, fill=FOREST, text_color=WHITE, width=None):
    size = 7.3
    w = width or (stringWidth(label, "Helvetica-Bold", size) + 15)
    c.setFillColor(fill)
    c.roundRect(x, y_top-15, w, 15, 7.5, stroke=0, fill=1)
    pdf_text(c, x+w/2, y_top-10.7, label.upper(), size, text_color, "Helvetica-Bold", align="center")
    return w


def pdf_header(c, page_num, title, eyebrow="DAYLIGHT LAB / DAY CREAM"):
    c.setFillColor(CREAM)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    # top left brand mark
    c.setFillColor(FOREST)
    c.circle(M+7, H-29, 7, stroke=0, fill=1)
    c.setFillColor(GOLD)
    c.circle(M+7, H-29, 2.5, stroke=0, fill=1)
    pdf_text(c, M+20, H-25, eyebrow, 7.4, FOREST, "Helvetica-Bold")
    pdf_text(c, M, H-53, title, 22, INK, "Helvetica-Bold")
    pdf_rule(c, M, H-64, W-M)
    # footer
    pdf_rule(c, M, 28, W-M)
    pdf_text(c, M, 16, "Fragrance-free home R&D prototype | Not a sunscreen or a medical acne treatment", 6.8, MUTED)
    pdf_text(c, W-M, 16, f"{page_num:02d}  /  08", 7, FOREST, "Helvetica-Bold", align="right")


def pdf_cover(c):
    c.setFillColor(CREAM)
    c.rect(0, 0, W, H, stroke=0, fill=1)
    # decorative pale circles
    c.setFillColor(SAGE)
    c.circle(508, 720, 160, stroke=0, fill=1)
    c.setFillColor(PEACH)
    c.circle(558, 318, 112, stroke=0, fill=1)
    c.setFillColor(MINT)
    c.circle(70, 85, 75, stroke=0, fill=1)
    # bespoke sun/drop mark
    c.setFillColor(FOREST)
    c.circle(472, 682, 58, stroke=0, fill=1)
    c.setFillColor(CREAM)
    c.circle(472, 682, 45, stroke=0, fill=1)
    c.setFillColor(GOLD)
    c.circle(472, 682, 28, stroke=0, fill=1)
    c.setFillColor(FOREST)
    c.circle(472, 682, 13, stroke=0, fill=1)
    # rays
    c.setStrokeColor(FOREST); c.setLineWidth(2)
    for a in range(0, 360, 45):
        from math import cos, sin
        x1 = 472 + cos(a*pi/180)*70; y1 = 682 + sin(a*pi/180)*70
        x2 = 472 + cos(a*pi/180)*84; y2 = 682 + sin(a*pi/180)*84
        c.line(x1, y1, x2, y2)
    pdf_badge(c, M, 742, "FORMULA + TUTORIAL", FOREST)
    pdf_text(c, M, 660, "DAYLIGHT", 42, INK, "Helvetica-Bold")
    pdf_text(c, M, 616, "DAY CREAM", 42, INK, "Helvetica-Bold")
    pdf_text(c, M, 578, "A practical, data-locked 100 g home R&D prototype", 14, FOREST, "Helvetica")
    pdf_text(c, M, 544, "Designed for fast visible hydration, smoother feel and healthy-looking glow - with a realistic long-term path for uneven tone.", 11, MUTED, maxw=390, leading=15)
    # goals panel
    pdf_card(c, M, 472, 408, 114, WHITE, LINE)
    pdf_text(c, M+18, 446, "THE DESIGN BRIEF", 7.5, FOREST, "Helvetica-Bold")
    pdf_text(c, M+18, 422, "4-DAY COSMETIC WIN", 14, INK, "Helvetica-Bold")
    pdf_text(c, M+18, 401, "Hydrated + fresh + smooth + comfortably glowy", 10.7, MUTED)
    pdf_rule(c, M+18, 383, M+388)
    pdf_text(c, M+18, 364, "pH 5.2-5.6  |  fragrance-free  |  1.00% preservative blend  |  no steroid / hydroquinone", 8.2, FOREST, "Helvetica-Bold")
    # intent / clock
    pdf_badge(c, M, 304, "REALISTIC EXPECTATION", GOLD, INK)
    pdf_text(c, M, 268, "Surface hydration can look better in days. Established pigmentation and acne biology need weeks, consistency and daytime sunscreen.", 13, INK, "Helvetica", maxw=460, leading=18)
    pdf_text(c, M, 178, "Prepared as an educational R&D document\nfor a single small, personal-use prototype.", 9, MUTED, leading=14)
    pdf_rule(c, M, 110, W-M)
    pdf_text(c, M, 86, "DAYLIGHT LAB", 8, FOREST, "Helvetica-Bold")
    pdf_text(c, M, 68, "Expert cosmetic formulation brief | India sourcing orientation | 01 Oct 2026", 8, MUTED)
    pdf_text(c, W-M, 68, "VERSION 1.0", 8, FOREST, "Helvetica-Bold", align="right")


def pdf_timeline_page(c):
    pdf_header(c, 2, "What can realistically change - and when")
    # top promise/guardrail
    pdf_card(c, M, 742, W-2*M, 78, MINT, SAGE)
    pdf_badge(c, M+17, 722, "FOUR-DAY FOCUS", FOREST)
    pdf_text(c, M+17, 691, "Humectants, light emollients and a smooth emulsion can quickly improve the way skin feels and reflects light. They cannot safely erase melanin or cure acne in four days.", 10, INK, maxw=W-2*M-34, leading=14)
    # four timeline cards
    y = 642
    colors = [SAGE, PEACH, MINT, WHITE]
    for i, (title, body) in enumerate(TIMELINE):
        h = 100 if i < 3 else 110
        pdf_card(c, M, y, W-2*M, h, colors[i], LINE)
        pdf_badge(c, M+16, y-17, title, FOREST if i in [0,2] else GOLD, INK if i==1 else WHITE)
        pdf_text(c, M+16, y-47, body, 10.1, INK, maxw=W-2*M-34, leading=14)
        y -= h + 12
    pdf_card(c, M, 148, W-2*M, 86, WHITE, LINE)
    pdf_text(c, M+16, 123, "USE ROUTINE", 8, FOREST, "Helvetica-Bold")
    pdf_text(c, M+16, 101, "Morning: apply a small amount to clean dry skin, then use a broad-spectrum SPF 30+ sunscreen. For pigment-prone skin, a tinted iron-oxide sunscreen is especially useful [3].", 9.4, INK, maxw=W-2*M-32, leading=13)
    pdf_text(c, M+16, 53, "This cream is not SPF-tested and must never be represented as sun protection.", 8.5, RED, "Helvetica-Bold")


def pdf_formula_table(c, rows, y_top, phase, title_note):
    x = M; width = W-2*M
    c.setFillColor(FOREST if phase == "A" else (GOLD if phase == "B" else TEAL))
    c.roundRect(x, y_top-26, width, 26, 10, stroke=0, fill=1)
    pdf_text(c, x+15, y_top-18, f"PHASE {phase}  |  {PHASES[phase].upper()}", 9, WHITE if phase != "B" else INK, "Helvetica-Bold")
    pdf_text(c, x+width-14, y_top-18, title_note, 7.6, WHITE if phase != "B" else INK, "Helvetica-Bold", align="right")
    y = y_top-38
    # Header
    cols = [x, x+164, x+393, x+445, x+497]
    c.setFillColor(SAGE); c.rect(x, y-18, width, 18, stroke=0, fill=1)
    headers = ["INGREDIENT", "WHAT IT DOES / WHY", "%", "GRAMS"]
    for hx, label in zip([cols[0]+8, cols[1]+8, cols[2]+10, cols[3]+10], headers):
        pdf_text(c, hx, y-12, label, 6.7, FOREST, "Helvetica-Bold")
    y -= 18
    for idx, (_, name, why, pct, _) in enumerate(rows):
        name_lines = pdf_wrap(name, "Helvetica-Bold", 8.3, 150)
        why_lines = pdf_wrap(why, "Helvetica", 7.7, 218)
        lines = max(len(name_lines), len(why_lines), 1)
        h = max(28, 10 + lines*10)
        c.setFillColor(WHITE if idx % 2 == 0 else MINT)
        c.rect(x, y-h, width, h, stroke=0, fill=1)
        yy = y-11
        for line in name_lines:
            pdf_text(c, x+8, yy, line, 8.3, INK, "Helvetica-Bold")
            yy -= 10
        yy = y-11
        for line in why_lines:
            pdf_text(c, x+172, yy, line, 7.7, MUTED, "Helvetica")
            yy -= 9.5
        pdf_text(c, x+425, y-h/2+2, f"{pct:.2f}", 8.3, INK, "Helvetica-Bold", align="right")
        pdf_text(c, x+486, y-h/2+2, f"{pct:.2f}", 8.3, INK, "Helvetica-Bold", align="right")
        c.setStrokeColor(LINE); c.setLineWidth(0.35); c.line(x, y-h, x+width, y-h)
        y -= h
    return y


def pdf_formula_pages(c):
    # Page 3
    pdf_header(c, 3, "The locked 100 g formula", "DAYLIGHT LAB / MASTER FORMULA")
    pdf_text(c, M, H-86, "Fragrance-free master formula. Percent equals grams for a 100 g batch. Total = 100.00 g.", 9.5, MUTED)
    A = [x for x in FORMULA if x[0] == "A"]
    B = [x for x in FORMULA if x[0] == "B"]
    y = pdf_formula_table(c, A, H-105, "A", "HEAT TO 70-75 C")
    y -= 19
    y = pdf_formula_table(c, B, y, "B", "HEAT TO 70-75 C")
    pdf_card(c, M, 78, W-2*M, 35, MINT, SAGE)
    pdf_text(c, M+12, 57, "Texture intention: light lotion-cream, quick-spreading, low-wax finish. The gel network + Olivem 1000 + cetearyl alcohol provide stability without a heavy feel.", 8.2, FOREST, "Helvetica", maxw=W-2*M-24, leading=11)

    # Page 4
    pdf_header(c, 4, "Cool-down actives, preservation & formula logic", "DAYLIGHT LAB / MASTER FORMULA")
    C_rows = [x for x in FORMULA if x[0] == "C"]
    y = pdf_formula_table(c, C_rows, H-87, "C", "ADD BELOW 40 C")
    # formula architecture cards
    y -= 17
    xw = (W-2*M-16)/3
    cards = [
        ("FAST LOOK", "4.00% glycerin + 3.00% propanediol + 1.00% betaine and light esters support quick hydration, slip and surface glow.", SAGE),
        ("LONGER PATH", "4.00% niacinamide + 2.00% NAG match concentrations used together in published pigmentation-appearance research [1].", PEACH),
        ("SAFETY CORE", "Target pH 5.2-5.6; 1.00% preservative blend; airless pack. Challenge testing is required before any sale [4,5].", MINT),
    ]
    for i, (heading, body, fill) in enumerate(cards):
        x = M+i*(xw+8)
        pdf_card(c, x, y, xw, 108, fill, LINE)
        pdf_text(c, x+12, y-23, heading, 8, FOREST, "Helvetica-Bold")
        pdf_text(c, x+12, y-43, body, 8.1, INK, maxw=xw-24, leading=11)
    # Optional fragrance note
    pdf_card(c, M, y-124, W-2*M, 49, WHITE, LINE)
    pdf_text(c, M+12, y-144, "SENSORY OPTION - NOT THE LOCKED FORMULA", 7.6, GOLD, "Helvetica-Bold")
    pdf_text(c, M+12, y-162, "Keep fragrance-free for the safest first prototype. If a scent is required, use 0.05% IFRA-compliant cosmetic fragrance and reduce total water by 0.05% (water becomes 68.64 g); patch test. Do not use essential oils as a shortcut.", 8.0, MUTED, maxw=W-2*M-24, leading=10.5)


def pdf_tutorial_page(c):
    pdf_header(c, 5, "Manufacturing tutorial - small, careful, repeatable")
    pdf_text(c, M, H-86, "Required: 0.01 g scale, calibrated pH meter, two heat-safe beakers, thermometer, silicone spatula, stick blender / mini homogeniser and opaque airless pump bottle.", 8.8, MUTED, maxw=W-2*M, leading=12)
    steps = [
        ("01  SANITIZE", "Clean bench, tools and pack. Spray 70% isopropyl alcohol; allow to air-dry fully. Wear clean gloves. Use only purified/deionized water - never tap water."),
        ("02  PHASE A", "Weigh the 43.69 g A-water. Pre-slurry 0.15 g xanthan into the 4.00 g glycerin. Add it to water with propanediol, betaine, EDTA and allantoin; stir until uniform."),
        ("03  PHASE B", "In the second beaker weigh CCT, dicaprylyl carbonate, squalane, Olivem 1000, cetearyl alcohol and tocopherol. Heat A and B separately to 70-75 C. Hold until B is fully melted and A is uniform (about 15-20 min)."),
        ("04  EMULSIFY", "With A and B within 3 C, slowly pour B into A. Blend 60-90 seconds in short pulses to minimise air. Then stir gently for 8-10 min. Do not whip."),
        ("05  COOL", "Stir gently while cooling. At 40 C or below, make sure phase C is a clear solution: dissolve citrate salts, niacinamide, NAG, panthenol and Zinc PCA in its 25.00 g water (gentle 35-40 C is allowed)."),
        ("06  FINISH", "Add phase C and the preservative blend below 40 C; mix 5 min. Check pH at 25 C after 10 min rest: target 5.2-5.6 (aim 5.4). Fill only when below 30 C."),
    ]
    y = H-118
    for i, (head, body) in enumerate(steps):
        h = 75 if i in [1,2,4] else 64
        fill = WHITE if i % 2 == 0 else MINT
        pdf_card(c, M, y, W-2*M, h, fill, LINE, 10)
        c.setFillColor(FOREST if i not in [2,5] else GOLD)
        c.circle(M+21, y-27, 11, stroke=0, fill=1)
        pdf_text(c, M+21, y-30, str(i+1), 8.2, WHITE if i not in [2,5] else INK, "Helvetica-Bold", align="center")
        pdf_text(c, M+42, y-20, head, 8.3, FOREST, "Helvetica-Bold")
        pdf_text(c, M+42, y-38, body, 8.4, INK, maxw=W-2*M-58, leading=11)
        y -= h+7
    pdf_card(c, M, 65, W-2*M, 42, PEACH, LINE)
    pdf_text(c, M+12, 47, "CRITICAL CONTROL", 7.4, RED, "Helvetica-Bold")
    pdf_text(c, M+12, 34, "Do not add the cool-down actives or preservative to a hot batch. If phase C is not fully dissolved, do not fill a gritty / unstable product.", 8.0, INK, maxw=W-2*M-24)


def pdf_control_page(c):
    pdf_header(c, 6, "pH, fill, stability & the small details that matter")
    # pH card
    pdf_card(c, M, H-88, W-2*M, 102, MINT, SAGE)
    pdf_badge(c, M+16, H-106, "MEASURED pH - DO NOT GUESS", FOREST)
    pdf_text(c, M+16, H-134, "The fixed citrate buffer is designed around pH 5.2-5.6 (aim 5.4 at 25 C), but raw-material lots change pH. Calibrate a meter with pH 4.01 and 7.00 buffers. If the first pilot is >5.6, titrate a 10% citric-acid solution one drop at a time, wait 10 min, recheck and record the amount; subtract the same mass from water in the next batch. If <5.2, do not use an unmeasured alkali - reformulate with a trained formulator.", 8.6, INK, maxw=W-2*M-32, leading=11.5)
    # fill left and right
    pdf_card(c, M, H-210, 250, 103, WHITE, LINE)
    pdf_text(c, M+14, H-234, "FILL + STORE", 8.3, FOREST, "Helvetica-Bold")
    fill_text = "Fill below 30 C into an opaque airless pump. Avoid jars: fingers raise contamination risk. Store sealed at 15-25 C, dry and away from sun. Because this home batch has no challenge test, make only 100 g and use within 4 weeks after opening."
    pdf_text(c, M+14, H-253, fill_text, 8.4, INK, maxw=222, leading=11)
    pdf_card(c, M+267, H-210, W-2*M-267, 103, WHITE, LINE)
    pdf_text(c, M+281, H-234, "HOME STABILITY SCREEN", 8.3, FOREST, "Helvetica-Bold")
    st_text = "Retain samples at room temperature, 4 C and 40 C for 4 weeks. Inspect weekly: separation, grains, odour, colour, pump function and pH. A pH shift >0.3, separation or odour change = fail. This is screening, not proof of shelf life."
    pdf_text(c, M+281, H-253, st_text, 8.4, INK, maxw=230, leading=11)
    # tricks
    pdf_text(c, M, H-336, "FORMULATION TRICKS THAT EARN THEIR PLACE", 10, INK, "Helvetica-Bold")
    tricks = [
        ("Keep phase temperatures aligned", "A/B within 3 C avoids shocky, weak emulsions."),
        ("Pre-slurry the xanthan", "Glycerin wetting prevents stubborn gum fish-eyes."),
        ("Use short blending pulses", "Enough shear for small droplets; less trapped air and foam."),
        ("Reserve cool-down for actives", "It protects panthenol and keeps the sensory profile clean."),
        ("Choose the airless pack", "Better user experience and lower in-use contamination than a jar."),
        ("Make small batches", "No challenge test means no claimed shelf life or sale."),
    ]
    y = H-360
    for i, (h, b) in enumerate(tricks):
        col = i % 2; row = i // 2
        x = M+col*267; yy = y-row*66
        pdf_card(c, x, yy, 250, 54, SAGE if i%2==0 else PEACH, LINE)
        pdf_text(c, x+12, yy-19, h, 8.1, FOREST, "Helvetica-Bold")
        pdf_text(c, x+12, yy-35, b, 8.0, INK, maxw=225, leading=10)
    # bug signs
    pdf_card(c, M, 136, W-2*M, 70, PEACH, LINE)
    pdf_text(c, M+14, 113, "STOP / DISCARD IF", 8.1, RED, "Helvetica-Bold")
    pdf_text(c, M+14, 91, "Mould, gas, unusual odour, colour change, new separation, watery thinning, gritty crystals, package swelling or unexpected irritation appear. Preservation must be proven by a suitable challenge test before commercial distribution [5].", 8.3, INK, maxw=W-2*M-28, leading=11)


def pdf_sourcing_cost_page(c):
    pdf_header(c, 7, "India shopping list & practical prototype cost")
    pdf_text(c, M, H-86, "Buy cosmetic raw materials with a current COA, SDS, batch number and expiry. Prices are approximate online orientation only, checked 01 Oct 2026; shipping, GST and pack size change them.", 8.6, MUTED, maxw=W-2*M, leading=11)
    # sourcing rows
    y = H-111
    for i, (material, grade, qty, note) in enumerate(SOURCING):
        h = 61 if i in [2,4,5,6] else 55
        pdf_card(c, M, y, W-2*M, h, WHITE if i%2==0 else MINT, LINE, 8)
        pdf_text(c, M+12, y-18, material, 8.2, INK, "Helvetica-Bold", maxw=150, leading=10)
        pdf_text(c, M+173, y-17, grade, 7.7, FOREST, "Helvetica", maxw=135, leading=9.5)
        pdf_text(c, M+319, y-17, qty, 7.7, INK, "Helvetica-Bold", maxw=50)
        pdf_text(c, M+375, y-16, note, 7.4, MUTED, "Helvetica", maxw=145, leading=9.2)
        y -= h+4
    # cost grid
    pdf_text(c, M, y-10, "BATCH COST ORIENTATION", 10, INK, "Helvetica-Bold")
    y -= 24
    widths = [95, 135, 135, 137]
    headers = ["BATCH", "RAW MATERIALS", "AIRLESS PACK + LABEL", "ESTIMATED TOTAL"]
    x = M
    for w, label in zip(widths, headers):
        c.setFillColor(FOREST); c.rect(x, y-22, w, 22, stroke=0, fill=1)
        pdf_text(c, x+8, y-14, label, 6.7, WHITE, "Helvetica-Bold")
        x += w
    y -= 22
    for idx, row in enumerate(COST_ROWS):
        x = M; h=25
        for w, val in zip(widths, row):
            c.setFillColor(WHITE if idx%2==0 else SAGE); c.rect(x,y-h,w,h,stroke=0,fill=1)
            pdf_text(c, x+8, y-16, val, 8, INK if x==M else FOREST, "Helvetica-Bold" if x==M else "Helvetica")
            x += w
        y -= h
    pdf_card(c, M, 97, W-2*M, 48, PEACH, LINE)
    pdf_text(c, M+12, 77, "IMPORTANT COST NOTE", 7.5, RED, "Helvetica-Bold")
    pdf_text(c, M+12, 62, "Minimum purchase for the first kit can be roughly Rs 3,500-6,000 before tools/shipping because actives and oils are sold in small packs. Do not cut cost by omitting preservative, pH measurement or clean packaging.", 8, INK, maxw=W-2*M-24, leading=10.5)


def pdf_safety_sources_page(c):
    pdf_header(c, 8, "Safety, boundaries & evidence")
    # safety cards
    left_x=M; right_x=M+267
    cards = [
        (left_x, H-86, "PATCH TEST", "Apply a rice-grain amount to the inner forearm or behind the ear for 2-3 days. Stop for burning, swelling, rash or persistent itching. Do not apply to broken / infected skin or the eye area."),
        (right_x, H-86, "PRESERVATION", "A 1.00% phenoxyethanol / ethylhexylglycerin blend is a sensible starting system, not proof. The actual product + pack needs preservative efficacy (challenge) testing before sale [4,5]."),
        (left_x, H-202, "DO NOT IMPROVISE", "Do not tip raw AHA/BHA, L-ascorbic acid, benzoyl peroxide, retinoids, essential oils or drug powders into this cream. They change pH, stability, irritation risk and preservation."),
        (right_x, H-202, "DAYTIME RULE", "Use this moisturiser under a broad-spectrum SPF 30+ sunscreen every day. For pigmentation, choose a tinted iron-oxide sunscreen where suitable [3]. This cream has no UV filters."),
    ]
    for x, y, title, body in cards:
        pdf_card(c, x, y, 250, 100, MINT if x==left_x else WHITE, LINE)
        pdf_text(c, x+13, y-22, title, 8.4, FOREST, "Helvetica-Bold")
        pdf_text(c, x+13, y-43, body, 8.45, INK, maxw=224, leading=11.5)
    pdf_card(c, M, H-322, W-2*M, 54, PEACH, LINE)
    pdf_text(c, M+14, H-344, "'NATURAL' IS NOT AUTOMATICALLY SAFER", 8.2, RED, "Helvetica-Bold")
    pdf_text(c, M+14, H-362, "Natural fragrance, citrus oil and unpreserved plant waters can sensitise skin or introduce microbes. Safety comes from identity, purity, concentration, preservation, testing and correct use - not a marketing word.", 8.3, INK, maxw=W-2*M-28, leading=11)
    # escalation
    pdf_card(c, M, H-394, W-2*M, 48, WHITE, LINE)
    pdf_text(c, M+14, H-414, "WHEN TO SEE A DERMATOLOGIST", 7.9, FOREST, "Helvetica-Bold")
    pdf_text(c, M+14, H-430, "Painful/cystic acne, scarring, sudden pigmentation, melasma, eczema, pregnancy-specific questions or a reaction need medical advice. This cosmetic prototype does not replace diagnosis or treatment.", 8.1, INK, maxw=W-2*M-28, leading=10)
    # references
    pdf_text(c, M, H-478, "KEY SOURCES", 10, INK, "Helvetica-Bold")
    yy = H-499
    for source in SOURCE_LINES:
        pdf_text(c, M+4, yy, source, 7.3, MUTED, "Helvetica", maxw=W-2*M-8, leading=9.2)
        yy -= 38 if len(source) > 170 else 28
    pdf_card(c, M, 64, W-2*M, 33, SAGE, LINE)
    pdf_text(c, M+12, 45, "DATA LOCK CHECK  |  Both deliverables are generated from one formula list: total 100.00 g, pH 5.2-5.6, A/B 70-75 C, C below 40 C, fill below 30 C, preservative blend 1.00%.", 7.6, FOREST, "Helvetica-Bold", maxw=W-2*M-24, leading=10)


def make_pdf():
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4, pageCompression=1)
    c.setTitle("Daylight Day Cream - Formula + Tutorial")
    c.setAuthor("Daylight Lab")
    c.setSubject("100 g fragrance-free day cream R&D prototype")
    pdf_cover(c); c.showPage()
    pdf_timeline_page(c); c.showPage()
    pdf_formula_pages(c); c.showPage()
    # formula pages function draws page 3 and page 4 in the same canvas; split correctly
    # The first page currently exists; it needs a page break between its two parts. Rebuild is handled below.
    c.save()

# The PDF needs an explicit split after page 3. Keep a separate builder so page numbering is deterministic.
def build_pdf():
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4, pageCompression=1)
    c.setTitle("Daylight Day Cream - Formula + Tutorial")
    c.setAuthor("Daylight Lab")
    c.setSubject("100 g fragrance-free day cream R&D prototype")
    pdf_cover(c); c.showPage()
    pdf_timeline_page(c); c.showPage()
    # Page 3 rendered alone
    pdf_header(c, 3, "The locked 100 g formula", "DAYLIGHT LAB / MASTER FORMULA")
    pdf_text(c, M, H-86, "Fragrance-free master formula. Percent equals grams for a 100 g batch. Total = 100.00 g.", 9.5, MUTED)
    A = [x for x in FORMULA if x[0] == "A"]
    B = [x for x in FORMULA if x[0] == "B"]
    y = pdf_formula_table(c, A, H-105, "A", "HEAT TO 70-75 C")
    y -= 19
    pdf_formula_table(c, B, y, "B", "HEAT TO 70-75 C")
    pdf_card(c, M, 78, W-2*M, 35, MINT, SAGE)
    pdf_text(c, M+12, 57, "Texture intention: light lotion-cream, quick-spreading, low-wax finish. The gel network + Olivem 1000 + cetearyl alcohol provide stability without a heavy feel.", 8.2, FOREST, "Helvetica", maxw=W-2*M-24, leading=11)
    c.showPage()
    # Page 4 rendered alone
    pdf_header(c, 4, "Cool-down actives, preservation & formula logic", "DAYLIGHT LAB / MASTER FORMULA")
    C_rows = [x for x in FORMULA if x[0] == "C"]
    y = pdf_formula_table(c, C_rows, H-87, "C", "ADD BELOW 40 C")
    y -= 17
    xw = (W-2*M-16)/3
    cards = [
        ("FAST LOOK", "4.00% glycerin + 3.00% propanediol + 1.00% betaine and light esters support quick hydration, slip and surface glow.", SAGE),
        ("LONGER PATH", "4.00% niacinamide + 2.00% NAG match concentrations used together in published pigmentation-appearance research [1].", PEACH),
        ("SAFETY CORE", "Target pH 5.2-5.6; 1.00% preservative blend; airless pack. Challenge testing is required before any sale [4,5].", MINT),
    ]
    for i, (heading, body, fill) in enumerate(cards):
        x = M+i*(xw+8)
        pdf_card(c, x, y, xw, 108, fill, LINE)
        pdf_text(c, x+12, y-23, heading, 8, FOREST, "Helvetica-Bold")
        pdf_text(c, x+12, y-43, body, 8.1, INK, maxw=xw-24, leading=11)
    pdf_card(c, M, y-124, W-2*M, 49, WHITE, LINE)
    pdf_text(c, M+12, y-144, "SENSORY OPTION - NOT THE LOCKED FORMULA", 7.6, GOLD, "Helvetica-Bold")
    pdf_text(c, M+12, y-162, "Keep fragrance-free for the safest first prototype. If a scent is required, use 0.05% IFRA-compliant cosmetic fragrance and reduce total water by 0.05% (water becomes 68.64 g); patch test. Do not use essential oils as a shortcut.", 8.0, MUTED, maxw=W-2*M-24, leading=10.5)
    c.showPage()
    pdf_tutorial_page(c); c.showPage()
    pdf_control_page(c); c.showPage()
    pdf_sourcing_cost_page(c); c.showPage()
    pdf_safety_sources_page(c); c.showPage()
    c.save()

# ---------------------------------------------------------------------------
# PowerPoint helpers
# ---------------------------------------------------------------------------
SLIDE_W = 13.333333
SLIDE_H = 7.5


def rgb(name):
    return RGBColor(*C[name])


def ppt_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color) if isinstance(color, str) else RGBColor(*color)
    if transparency:
        shape.fill.transparency = transparency
    shape.line.fill.background()


def ppt_line(shape, color="line", width=1):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def ppt_text(slide, x, y, w, h, text, size=12, color="ink", bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04, fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = rgb(color) if isinstance(color, str) else RGBColor(*color)
    if fit:
        try: tf.fit_text(font_family=font, max_size=Pt(size))
        except Exception: pass
    return box


def ppt_card(slide, x, y, w, h, fill="white", line="line", radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    ppt_fill(shape, fill)
    if line:
        ppt_line(shape, line, 0.7)
    return shape


def ppt_circle(slide, x, y, d, fill="forest", line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    ppt_fill(shape, fill)
    if line: ppt_line(shape, line, .7)
    return shape


def ppt_header(slide, page, kicker="DAYLIGHT LAB / DAY CREAM"):
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = rgb("cream")
    ppt_circle(slide, .38, .23, .13, "forest")
    ppt_circle(slide, .423, .273, .045, "gold")
    ppt_text(slide, .58, .2, 4.5, .25, kicker, 6.8, "forest", True)
    # Page label
    ppt_text(slide, 11.95, .20, 1.0, .22, f"{page:02d}  /  11", 6.8, "forest", True, align=PP_ALIGN.RIGHT)
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.38), Inches(.53), Inches(12.57), Inches(.012))
    ppt_fill(ln, "line")
    # footer
    fline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.38), Inches(7.14), Inches(12.57), Inches(.01))
    ppt_fill(fline, "line")
    ppt_text(slide, .38, 7.18, 7.9, .18, "Fragrance-free home R&D prototype | Not a sunscreen or medical acne treatment", 5.4, "muted")


def ppt_badge(slide, x, y, text, fill="forest", color="white", w=None):
    w = w or max(1.08, 0.055*len(text)+.23)
    card = ppt_card(slide, x, y, w, .24, fill, line=None)
    ppt_text(slide, x+.04, y+.038, w-.08, .13, text.upper(), 5.4, color, True, align=PP_ALIGN.CENTER, margin=0)
    return card


def ppt_title(slide, title, sub=None):
    ppt_text(slide, .38, .72, 12.0, .4, title, 23, "ink", True)
    if sub:
        ppt_text(slide, .40, 1.16, 11.9, .28, sub, 8.4, "muted")


def ppt_table(slide, rows, x, y, col_widths, row_heights=None, header=None, font_size=6.4, include_phase=False):
    # rows list, with `header` optional.
    width = sum(col_widths)
    cur_y = y
    if header:
        ppt_card(slide, x, cur_y, width, .28, "forest", line=None, radius=False)
        xx=x
        for text, cw in zip(header, col_widths):
            ppt_text(slide, xx+.05, cur_y+.075, cw-.10, .11, text, 5.4, "white", True, margin=0)
            xx += cw
        cur_y += .28
    for idx, row in enumerate(rows):
        h = row_heights[idx] if row_heights else .37
        ppt_card(slide, x, cur_y, width, h, "white" if idx%2==0 else "mint", line="line", radius=False)
        xx=x
        for j, (text, cw) in enumerate(zip(row, col_widths)):
            col = "ink" if j != 1 else "muted"
            bold = (j == 0) or (j >= len(row)-2)
            fs = font_size if j != 1 else max(font_size-.15, 5.7)
            ppt_text(slide, xx+.055, cur_y+.05, cw-.11, h-.09, str(text), fs, col, bold, margin=0, fit=True)
            xx += cw
        cur_y += h
    return cur_y


def p_formula_rows(phase):
    return [[name, why, f"{pct:.2f}", f"{pct:.2f}"] for p, name, why, pct, _ in FORMULA if p == phase]


def make_ppt():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W); prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    # Slide 1 cover
    s = prs.slides.add_slide(blank)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb("cream")
    # decorative rings, right
    ppt_circle(s, 9.45, .55, 3.1, "sage", None)
    ppt_circle(s, 9.72, .82, 2.56, "cream", None)
    ppt_circle(s, 10.12, 1.22, 1.76, "gold", None)
    ppt_circle(s, 10.53, 1.63, .94, "forest", None)
    # lower accent
    ppt_circle(s, -1.1, 5.75, 2.7, "mint", None)
    ppt_badge(s, .52, .72, "FORMULA + TUTORIAL", "forest")
    ppt_text(s, .52, 1.42, 7.8, .6, "DAYLIGHT", 34, "ink", True)
    ppt_text(s, .52, 2.03, 7.8, .6, "DAY CREAM", 34, "ink", True)
    ppt_text(s, .55, 2.74, 7.7, .38, "A practical, data-locked 100 g home R&D prototype", 13, "forest")
    ppt_text(s, .55, 3.25, 7.25, .7, "Fast visible hydration, smoother feel and healthy-looking glow - with a realistic long-term path for uneven tone.", 13, "muted")
    ppt_card(s, .52, 4.45, 8.25, 1.25, "white", "line")
    ppt_text(s, .8, 4.73, 3.0, .16, "THE FOUR-DAY BRIEF", 7.2, "forest", True)
    ppt_text(s, .8, 5.00, 7.4, .26, "Hydrated  |  Fresh  |  Smooth  |  Comfortably glowy", 14.2, "ink", True)
    ppt_text(s, .8, 5.41, 7.5, .16, "pH 5.2-5.6  |  fragrance-free  |  no steroid / hydroquinone", 7.3, "muted")
    ppt_text(s, .55, 6.54, 5.8, .2, "DAYLIGHT LAB  |  VERSION 1.0  |  01 OCT 2026", 6.8, "forest", True)

    # Slide 2: goal and realism
    s = prs.slides.add_slide(blank); ppt_header(s, 2); ppt_title(s, "What can realistically change - and when", "Optimise surface optics now; build a longer-term routine without promising medical outcomes.")
    ppt_card(s, .42, 1.62, 12.48, .64, "mint", "sage")
    ppt_badge(s, .66, 1.80, "FOUR-DAY FOCUS", "forest")
    ppt_text(s, 2.15, 1.77, 10.3, .30, "Humectants + light emollients can improve hydration, feel and light reflection quickly. They do not erase melanin or cure acne in four days.", 9.1, "ink")
    positions=[(.43,2.57, "DAY 1-4", "Hydrated, smoother, fresher-looking; more even light reflection. Not pigment removal or acne cure.", "sage"),
               (6.72,2.57,"1 WEEK", "Less tightness/flaking and a more comfortable feel if the formula suits the skin. Breakouts may be unchanged.", "peach"),
               (.43,4.27,"2-4 WEEKS", "Gradual improvement in overall tone appearance may begin. Results vary with skin, routine and UV exposure.", "mint"),
               (6.72,4.27,"6-12 WEEKS", "A longer-term target for uneven pigmentation. Daily sunscreen is essential; persistent acne/melasma needs a dermatologist.", "white")]
    for x,y,head,body,fill in positions:
        ppt_card(s,x,y,6.15,1.35,fill,"line")
        ppt_badge(s,x+.24,y+.22,head,"forest" if head in ["DAY 1-4","2-4 WEEKS"] else "gold", "white" if head in ["DAY 1-4","2-4 WEEKS"] else "ink")
        ppt_text(s,x+.24,y+.65,5.67,.46,body,8.2,"ink")
    ppt_card(s,.43,6.15,12.45,.58,"white","line")
    ppt_text(s,.67,6.35,11.9,.17,"ROUTINE: apply the cream, then a broad-spectrum SPF 30+ sunscreen. Tinted iron-oxide sunscreen is useful for dark spots [3]. This cream is not SPF-tested.",7.8,"forest",True)

    # Slide 3: formula architecture
    s = prs.slides.add_slide(blank); ppt_header(s,3); ppt_title(s, "Design architecture", "A compact emulsion focused on immediate cosmetic benefits first, with well-established longer-pathway actives.")
    # architecture horizontal flow
    flow=[("WATER + HUMECTANTS","Glycerin 4.00%\nPropanediol 3.00%\nBetaine 1.00%","Fast hydration + plump surface look","sage"),
          ("LIGHT EMOLLIENTS","CCT 4.00%\nDicaprylyl carbonate 3.00%\nSqualane 2.00%","Slip, softness, lower-wax feel","peach"),
          ("TONE PATHWAY","Niacinamide 4.00%\nNAG 2.00%","Published combination levels [1]","mint"),
          ("CONTROL SYSTEM","pH 5.2-5.6\nPreservative blend 1.00%","Comfort + controlled prototype","white")]
    for i,(head,chem,reason,fill) in enumerate(flow):
        x=.45+i*3.12
        ppt_card(s,x,1.72,2.78,2.55,fill,"line")
        ppt_circle(s,x+.95,1.94,.85,"forest" if i!=1 else "gold")
        ppt_text(s,x+.95,2.21,.85,.18,str(i+1),10,"white" if i!=1 else "ink",True,align=PP_ALIGN.CENTER)
        ppt_text(s,x+.20,3.0,2.38,.24,head,7.2,"forest",True,align=PP_ALIGN.CENTER)
        ppt_text(s,x+.25,3.40,2.28,.51,chem,8.2,"ink",True,align=PP_ALIGN.CENTER)
        ppt_text(s,x+.20,4.01,2.38,.27,reason,6.9,"muted",align=PP_ALIGN.CENTER)
        if i<3:
            arrow=s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x+2.83), Inches(2.75), Inches(.23), Inches(.33)); ppt_fill(arrow,"forest")
    ppt_card(s,.45,4.72,12.45,1.33,"white","line")
    ppt_text(s,.72,4.99,2.0,.19,"WHY THIS IS FAST",8,"forest",True)
    ppt_text(s,.72,5.31,5.38,.45,"Humectants bind water in the outer skin. Light emollients smooth the micro-surface. Together they can create the fastest safe “glow” effect.",8.4,"ink")
    ppt_text(s,6.68,4.99,2.6,.19,"WHY THIS IS NOT A MIRACLE",8,"forest",True)
    ppt_text(s,6.68,5.31,5.52,.45,"Pigmentation depends on skin turnover and UV/visible-light exposure; pimples have several causes. Do not make cure claims.",8.4,"ink")
    ppt_badge(s,.52,6.31,"FRAGRANCE-FREE DEFAULT","forest")
    ppt_text(s,2.17,6.34,9.7,.17,"Optional sensory variant: 0.05% IFRA-compliant cosmetic fragrance replaces 0.05% water (68.64 g total water).",7.1,"muted")

    # Slide 4 Formula A
    s=prs.slides.add_slide(blank); ppt_header(s,4,"DAYLIGHT LAB / MASTER FORMULA"); ppt_title(s,"Locked 100 g formula - Phase A", "Fragrance-free. Percent equals grams for this 100 g batch. Phase A heat range: 70-75 C.")
    ppt_badge(s,.43,1.49,"PHASE A / WATER PHASE","forest")
    arows=p_formula_rows("A")
    ppt_table(s,arows,.43,1.85,[3.2,6.35,1.1,1.1],header=["INGREDIENT","WHAT IT DOES / WHY","%","GRAMS"],font_size=6.4)
    ppt_card(s,.43,5.47,12.45,.82,"mint","sage")
    ppt_text(s,.70,5.70,11.9,.20,"PHASE A PROCESS",7.3,"forest",True)
    ppt_text(s,.70,5.98,11.6,.18,"Pre-slurry xanthan in glycerin; combine with A-water, propanediol, betaine, EDTA and allantoin. Heat and hold the complete A phase at 70-75 C until uniform.",7.6,"ink")
    ppt_card(s,.43,6.46,12.45,.36,"white","line")
    ppt_text(s,.67,6.58,11.9,.12,"A-water = 43.69 g; cool-down water is listed separately in Phase C. Formula total is maintained at 100.00 g.",6.8,"muted")

    # Slide 5 Formula B&C
    s=prs.slides.add_slide(blank); ppt_header(s,5,"DAYLIGHT LAB / MASTER FORMULA"); ppt_title(s,"Locked 100 g formula - Phases B & C", "Phase B heat range: 70-75 C. Phase C is added below 40 C.")
    ppt_badge(s,.43,1.45,"PHASE B / OIL PHASE","gold","ink")
    brows=p_formula_rows("B")
    ppt_table(s,brows,.43,1.78,[3.2,6.35,1.1,1.1],row_heights=[.31]*len(brows),header=["INGREDIENT","WHAT IT DOES / WHY","%","GRAMS"],font_size=5.8)
    ppt_badge(s,.43,4.10,"PHASE C / COOL-DOWN ACTIVES + PRESERVATION","teal")
    # Compact visual copy; ingredient identity and exact values still come from FORMULA.
    crows=[[name, why.replace("Brightening/barrier-support active; selected at a studied 4% level.", "Tone + barrier support.").replace("Tone-evenness support; pairs well with niacinamide.", "Tone-evenness support.").replace("Humectant/barrier-support active for skin comfort.", "Hydration + skin comfort.").replace("Blemish-prone/oil-balance support; not an acne medicine.", "Blemish/oil support; not medicine.").replace("Part of a gentle citrate buffer for the target pH.", "Citrate buffer.").replace("Part of the citrate buffer; supports pH control.", "Citrate buffer.").replace("Broad preservative blend (e.g., Euxyl PE 9010).", "Preservative blend."), f"{pct:.2f}", f"{pct:.2f}"] for p,name,why,pct,_ in FORMULA if p=="C"]
    ppt_table(s,crows,.43,4.42,[3.2,6.35,1.1,1.1],row_heights=[.25]*len(crows),header=["INGREDIENT","WHAT IT DOES / WHY","%","GRAMS"],font_size=5.1)
    ppt_text(s,.45,6.94,12.0,.13,"Data lock: total 100.00 g | target pH 5.2-5.6 (aim 5.4 at 25 C) | Phenoxyethanol (and) Ethylhexylglycerin blend = 1.00%.",6.3,"forest",True)

    # Slide 6 tutorial flow
    s=prs.slides.add_slide(blank); ppt_header(s,6); ppt_title(s,"Manufacturing flow", "Use a 0.01 g scale, calibrated pH meter, thermometer, two heat-safe beakers, stick blender, gloves and opaque airless pump.")
    flow_steps=[("1", "SANITIZE", "70% IPA on tools and pack; air-dry. Purified water only."),
                ("2", "PHASE A", "Pre-slurry xanthan in glycerin. Combine A and heat to 70-75 C."),
                ("3", "PHASE B", "Combine oils + emulsifier. Heat to 70-75 C until melted."),
                ("4", "EMULSIFY", "A/B within 3 C. Add B into A; blend 60-90 sec in short pulses."),
                ("5", "COOL + C", "Cool below 40 C. Dissolve C actives in 25.00 g C-water; add with preservative."),
                ("6", "pH + FILL", "Mix 5 min; check pH after 10 min. Fill below 30 C into an airless pump.")]
    for i,(n,head,body) in enumerate(flow_steps):
        row=i//3; col=i%3; x=.45+col*4.14; y=1.66+row*2.18
        ppt_card(s,x,y,3.76,1.65,"white" if i%2==0 else "mint","line")
        ppt_circle(s,x+.22,y+.22,.46,"forest" if i not in [2,5] else "gold")
        ppt_text(s,x+.22,y+.36,.46,.13,n,7.6,"white" if i not in [2,5] else "ink",True,align=PP_ALIGN.CENTER)
        ppt_text(s,x+.80,y+.24,2.65,.18,head,8.3,"forest",True)
        ppt_text(s,x+.22,y+.86,3.25,.40,body,7.5,"ink")
        if col<2:
            arr=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x+3.84),Inches(y+.68),Inches(.22),Inches(.28)); ppt_fill(arr,"forest")
    ppt_card(s,.45,6.20,12.43,.58,"peach","line")
    ppt_text(s,.70,6.39,11.93,.17,"NON-NEGOTIABLE: do not heat phase C actives or preservative; do not fill a batch that is gritty, separated, incorrectly pH-adjusted or visibly aerated.",7.4,"ink",True)

    # Slide 7 controls
    s=prs.slides.add_slide(blank); ppt_header(s,7); ppt_title(s,"Critical controls: pH, preservation & stability", "A beautiful prototype still needs measurement, packaging discipline and a no-sale boundary.")
    controls=[("pH 5.2-5.6", "Measure at 25 C after 10 min rest. Calibrate meter pH 4.01 / 7.00. Fixed citrate buffer helps; raw lots can differ."),
              ("1.00% preservative blend", "Phenoxyethanol (and) Ethylhexylglycerin. It is a starting system, not a challenge-test result [4,5]."),
              ("AIRLESS, 15-25 C", "Fill below 30 C. Use opaque airless pack, not a finger-dip jar. Store dry, away from sun."),
              ("4-WEEK HOME LIMIT", "No challenge test = make only 100 g and use within 4 weeks after opening. No sale / no shelf-life claim."),
              ("4-WEEK SCREEN", "Hold samples at room temperature, 4 C and 40 C. Check weekly: pH, odour, colour, separation and pump."),
              ("FAIL / DISCARD", "Mould, gas, odd odour, colour shift, watery thinning, separation, crystals, swelling or unexpected irritation.")]
    for i,(head,body) in enumerate(controls):
        row=i//3; col=i%3; x=.45+col*4.14; y=1.66+row*2.15
        fill=["sage","peach","mint","white","white","peach"][i]
        ppt_card(s,x,y,3.76,1.63,fill,"line")
        ppt_text(s,x+.23,y+.24,3.25,.19,head,8.3,"forest" if i!=5 else "red",True)
        ppt_text(s,x+.23,y+.63,3.23,.60,body,7.45,"ink")
    ppt_card(s,.45,6.08,12.43,.58,"mint","sage")
    ppt_text(s,.70,6.27,11.93,.17,"IF pH IS HIGH: titrate 10% citric-acid solution one drop at a time, rest 10 min, recheck and record. Rebalance water by that amount in the next batch. IF LOW: do not add unmeasured alkali.",7.0,"forest",True)

    # Slide 8 sourcing
    s=prs.slides.add_slide(blank); ppt_header(s,8); ppt_title(s,"India sourcing checklist", "Purchase exact INCI materials only from cosmetic raw-material suppliers that provide COA, SDS, batch and expiry.")
    # shortened sourcing list visual
    source_cards=[("ACTIVES","Niacinamide >=99%, NAG >=98%, D-Panthenol >=98%, Zinc PCA","ASES lists niacinamide 99.5% / 100 g around Rs 240. Cosmesi Global lists Zinc PCA; 100 g around Rs 750. Confirm current COA / stock.","sage"),
                  ("EMULSION CORE","Olivem 1000, cetearyl alcohol, CCT, dicaprylyl carbonate, squalane","Euroasia Cosmetics and other Indian cosmetic raw suppliers list these categories. Verify exact INCI and cosmetic grade.","peach"),
                  ("SUPPORT CORE","DI water, USP glycerin, 1,3-propanediol, betaine, allantoin, EDTA, transparent xanthan, citrate salts","Buy 25-100 g packs with documents. Do not use unknown food or pharmacy powders as substitutes.","mint"),
                  ("PRESERVATIVE","Phenoxyethanol (and) Ethylhexylglycerin (e.g., Euxyl PE 9010)","Use a verified blend with correct INCI. Follow its technical sheet. Do not substitute a finished medicine or omit preservation.","white")]
    for i,(head,items,note,fill) in enumerate(source_cards):
        row=i//2;col=i%2;x=.45+col*6.27;y=1.66+row*2.10
        ppt_card(s,x,y,6.0,1.72,fill,"line")
        ppt_badge(s,x+.22,y+.21,head,"forest" if i!=1 else "gold","white" if i!=1 else "ink")
        ppt_text(s,x+.24,y+.60,5.5,.30,items,7.5,"ink",True)
        ppt_text(s,x+.24,y+1.08,5.42,.38,note,6.65,"muted")
    ppt_card(s,.45,6.12,12.43,.57,"peach","line")
    ppt_text(s,.69,6.31,11.93,.17,"FIRST KIT BUDGET: roughly Rs 3,500-6,000 before tools/shipping because of small-pack MOQs. Never save money by omitting the preservative, pH meter or clean airless pack.",7.35,"ink",True)

    # Slide 9 cost
    s=prs.slides.add_slide(blank); ppt_header(s,9); ppt_title(s,"Cost orientation", "Indicative India small-batch costs. Includes ingredient allowance and airless pack / label, but excludes labour, tools, shipping, GST and testing.")
    headers=["BATCH","RAW MATERIALS","AIRLESS PACK + LABEL","ESTIMATED TOTAL"]
    rows=COST_ROWS
    ppt_table(s,rows,.88,1.80,[1.55,2.9,3.25,3.05],row_heights=[.62,.62,.62],header=headers,font_size=10)
    ppt_card(s,.88,4.16,10.75,1.23,"mint","sage")
    ppt_text(s,1.14,4.43,10.2,.18,"HOW TO READ THIS",8.2,"forest",True)
    ppt_text(s,1.14,4.76,9.95,.36,"Costs decrease at 1 kg because pack price falls. The first 100 g test is not the same as commercial COGS: preservation challenge test, stability, compatibility, regulatory review and production controls are not included.",8.1,"ink")
    # visual stacked bar, decorative not numerical total
    ppt_text(s,.88,5.83,4.0,.19,"SPEND PRIORITY",7.6,"forest",True)
    segments=[("Identity + COA",2.25,"forest"),("Actives",2.8,"gold"),("Pack",1.7,"teal"),("Testing",3.4,"sage")]
    xx=.88
    for label,wid,fill in segments:
        ppt_card(s,xx,6.17,wid,.34,fill,None,False)
        ppt_text(s,xx+.05,6.28,wid-.1,.1,label,5.5,"white" if fill in ["forest","teal"] else "ink",True,align=PP_ALIGN.CENTER)
        xx+=wid+.05
    ppt_text(s,.88,6.72,10.7,.12,"For a product to sell: professional preservation challenge testing and stability testing are required expenditures, not optional marketing extras.",6.8,"muted")

    # Slide 10 safety
    s=prs.slides.add_slide(blank); ppt_header(s,10); ppt_title(s,"Safety and routine boundaries", "The goal is visible cosmetic improvement without unsafe shortcuts.")
    safety=[("PATCH TEST", "Use a rice-grain amount for 2-3 days on inner forearm / behind ear. Stop for burning, swelling, rash or persistent itch."),
            ("DO NOT IMPROVISE", "Do not add raw acids, L-ascorbic acid, benzoyl peroxide, retinoids, essential oils or drug powders into this base."),
            ("SUNSCREEN REQUIRED", "Morning cream first; broad-spectrum SPF 30+ sunscreen second. Tinted iron oxides help pigment-prone skin [3]."),
            ("NATURAL IS NOT A SAFETY CLAIM", "Natural fragrance and plant waters can sensitise or introduce microbes. Safety is identity + dose + preservation + testing."),
            ("DOCTOR FIRST WHEN NEEDED", "Cystic/painful acne, scarring, sudden pigmentation, melasma, eczema or reactions require medical advice."),
            ("NOT FOR SALE", "This unchallenged 100 g prototype is personal R&D only. Product sale needs formal safety, stability, microbiology and regulatory work.")]
    for i,(head,body) in enumerate(safety):
        row=i//3; col=i%3; x=.45+col*4.14; y=1.65+row*2.14
        ppt_card(s,x,y,3.76,1.62,"white" if i%2 else "mint","line")
        ppt_circle(s,x+.22,y+.23,.42,"forest" if i not in [3,5] else "gold")
        ppt_text(s,x+.22,y+.36,.42,.1,"!",7.2,"white" if i not in [3,5] else "ink",True,align=PP_ALIGN.CENTER)
        ppt_text(s,x+.79,y+.25,2.72,.17,head,7.6,"forest" if i!=4 else "red",True)
        ppt_text(s,x+.22,y+.78,3.22,.48,body,7.25,"ink")
    ppt_card(s,.45,6.10,12.43,.58,"peach","line")
    ppt_text(s,.69,6.29,11.93,.18,"SAFE POSITIONING: moisturising / glow / fresh-looking / supports an even-looking tone. Do not claim to treat acne, melasma or pigmentation disease.",7.35,"ink",True)

    # Slide 11 sources / lock
    s=prs.slides.add_slide(blank); ppt_header(s,11); ppt_title(s,"Evidence base & data lock", "The PDF and this deck are generated from the same master formula data list.")
    ppt_card(s,.45,1.62,12.43,.78,"sage","line")
    ppt_text(s,.70,1.85,11.93,.18,"DATA LOCK  |  Total 100.00 g  |  pH 5.2-5.6 (aim 5.4)  |  A/B 70-75 C  |  C below 40 C  |  fill below 30 C  |  preservative blend 1.00%",7.25,"forest",True)
    ppt_text(s,.70,2.12,11.93,.13,"Formula is fragrance-free. Optional scented variant only: 0.05% IFRA-compliant fragrance replaces 0.05% water; total water 68.64 g.",6.5,"muted")
    y=2.72
    for source in SOURCE_LINES:
        ppt_card(s,.45,y,12.43,.62,"white","line")
        ppt_text(s,.66,y+.12,12.0,.36,source,6.15,"ink")
        y+=.72
    ppt_card(s,.45,6.40,12.43,.34,"mint","sage")
    ppt_text(s,.68,6.51,11.95,.11,"Source claims are used for formulation rationale, not as a performance guarantee. Current supplier price examples are orientation only; verify at checkout.",6.2,"forest",True)

    # metadata
    prs.core_properties.title = "Daylight Day Cream - Formula Presentation"
    prs.core_properties.subject = "100 g fragrance-free day cream R&D prototype"
    prs.core_properties.author = "Daylight Lab"
    prs.core_properties.keywords = "cosmetic formulation, niacinamide, N-acetyl glucosamine, day cream"
    prs.save(str(PPTX_PATH))

# ---------------------------------------------------------------------------
# Cross checks and CLI
# ---------------------------------------------------------------------------
def formula_report():
    lines = []
    for p, name, why, pct, group in FORMULA:
        lines.append(f"{p}\t{name}\t{pct:.2f}\t{pct:.2f}")
    return "\n".join(lines) + f"\nTOTAL\t{sum(x[3] for x in FORMULA):.2f}\nPH\t{PH_TARGET}\nHEAT\t{HEAT_RANGE}\nCOOLDOWN\t{COOL_ADD}\n"


def check_pptx_data():
    # inspect text-bearing XML for all canonical names and values
    import zipfile
    with zipfile.ZipFile(PPTX_PATH) as z:
        text = "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
    missing = []
    for _, name, _, pct, _ in FORMULA:
        # XML text can escape ampersand; all names are present exactly except ampersand XML escaping.
        token = name.replace("&", "&amp;")
        if token not in text:
            missing.append(name)
        if f"{pct:.2f}" not in text:
            missing.append(f"value {pct:.2f}")
    # Key numerical locks must occur.
    for token in ["100.00", "5.2-5.6", "70-75", "below 40", "below 30", "1.00"]:
        if token not in text: missing.append(token)
    if missing:
        raise RuntimeError("PPTX data lock check failed: " + ", ".join(missing))


def check_pdf_exists():
    if not PDF_PATH.exists() or PDF_PATH.stat().st_size < 15_000:
        raise RuntimeError("PDF was not created correctly")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_pdf()
    make_ppt()
    check_pdf_exists()
    check_pptx_data()
    (OUT / "formula_data_lock.tsv").write_text(formula_report(), encoding="utf-8")
    print(f"Created {PDF_PATH.name} ({PDF_PATH.stat().st_size:,} bytes)")
    print(f"Created {PPTX_PATH.name} ({PPTX_PATH.stat().st_size:,} bytes)")
    print("Cross-check: canonical formula total = 100.00 g; PPTX master data verified.")
